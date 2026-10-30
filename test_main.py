from decimal import Decimal

from main import parse_price, normalize_headers


def test_valid_price():
    assert parse_price(1500) == Decimal("1500")
    assert parse_price("1450,50") == Decimal("1450.50")


def test_invalid_price():
    assert parse_price(None) is None
    assert parse_price(-100) is None
    assert parse_price("abc") is None
    assert parse_price(0) is None


def test_normalize_headers():
    headers = ["SKU", "Товар", "Стоимость"]

    result = normalize_headers(headers)

    assert result == {
        "article": 0,
        "name": 1,
        "price": 2,
    }
from openpyxl import Workbook, load_workbook
import main


def test_full_excel_processing(tmp_path, monkeypatch):
    # Создаём временные папки
    input_dir = tmp_path / "input"
    output_dir = tmp_path / "output"
    input_dir.mkdir()

    # Создаём тестовый Excel-файл
    wb = Workbook()
    ws = wb.active
    ws.append(["Артикул", "Название", "Цена"])
    ws.append(["A101", "Мышь", 1500])
    ws.append(["A101", "Мышь", 1300])
    ws.append(["A102", "Клавиатура", None])
    wb.save(input_dir / "supplier.xlsx")

    # Направляем программу во временные папки
    monkeypatch.setattr(main, "INPUT_DIR", input_dir)
    monkeypatch.setattr(main, "OUTPUT_DIR", output_dir)

    # Запускаем основную обработку
    products, errors = main.process_files()
    main.save_result(products, errors)

    # Проверяем созданный Excel-файл
    result = load_workbook(output_dir / "result.xlsx")

    try:
        product_rows = list(
            result["Товары"].iter_rows(values_only=True)
        )
        error_rows = list(
            result["Ошибки"].iter_rows(values_only=True)
        )

        assert len(product_rows) == 2  # Заголовок + 1 товар
        assert product_rows[1] == ("A101", "Мышь", 1300)

        assert len(error_rows) == 2  # Заголовок + 1 ошибка
        assert error_rows[1][2] == "A102"
    finally:
        result.close()

def test_corrupted_excel_file(tmp_path, monkeypatch):
    input_dir = tmp_path / "input"
    input_dir.mkdir()

    # Создаём файл с расширением xlsx,
    # но с некорректным содержимым
    corrupted_file = input_dir / "broken.xlsx"
    corrupted_file.write_text("Это не Excel-файл")

    monkeypatch.setattr(main, "INPUT_DIR", input_dir)

    products, errors = main.process_files()

    assert products == {}
    assert len(errors) == 1
    assert errors[0][0] == "broken.xlsx"
    assert "Не удалось открыть файл" in errors[0][3]

def test_conflicting_product_names(tmp_path, monkeypatch):
    input_dir = tmp_path / "input"
    input_dir.mkdir()

    wb = Workbook()
    ws = wb.active
    ws.append(["Артикул", "Название", "Цена"])
    ws.append(["A101", "Мышь Logitech", 1500])
    ws.append(["A101", "Клавиатура Defender", 900])
    wb.save(input_dir / "supplier.xlsx")

    monkeypatch.setattr(main, "INPUT_DIR", input_dir)
    

    products, errors = main.process_files()

    assert any(
        "Конфликт названий" in error[3]
        for error in errors
    )

    assert "A101" not in products


def test_conflict_across_three_files(tmp_path, monkeypatch):
    input_dir = tmp_path / "input"
    input_dir.mkdir()

    suppliers = [
        ("supplier_1.xlsx", "Мышь Logitech", 1500),
        ("supplier_2.xlsx", "Клавиатура Defender", 900),
        ("supplier_3.xlsx", "Мышь Logitech", 1400),
    ]

    for filename, name, price in suppliers:
        wb = Workbook()
        ws = wb.active

        ws.append(["Артикул", "Название", "Цена"])
        ws.append(["A101", name, price])

        wb.save(input_dir / filename)
        wb.close()

    monkeypatch.setattr(main, "INPUT_DIR", input_dir)

    products, errors = main.process_files()

    assert "A101" not in products

    assert any(
        "Конфликт названий" in error[3]
        for error in errors
    )
