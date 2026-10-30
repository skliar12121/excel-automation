
from pathlib import Path
from decimal import Decimal, InvalidOperation
from openpyxl import Workbook, load_workbook

INPUT_DIR = Path("input")
OUTPUT_DIR = Path("output")

COLUMN_NAMES = {
    "Артикул": "article",
    "SKU": "article",
    "Код": "article",
    "Название": "name",
    "Товар": "name",
    "Наименование": "name",
    "Цена": "price",
    "Стоимость": "price",
    "Цена, руб.": "price",
}


def normalize_headers(headers):
    """Приводит названия столбцов к единому формату."""
    return {
        COLUMN_NAMES[str(header).strip()]: index
        for index, header in enumerate(headers)
        if str(header).strip() in COLUMN_NAMES
    }


def parse_price(value):
    """Проверяет цену и преобразует её в число."""
    if value is None or isinstance(value, bool):
        return None

    try:
        price = Decimal(str(value).strip().replace(",", "."))
        if not price.is_finite() or price <= 0:
            return None
        return price
    except (InvalidOperation, ValueError):
        return None


def process_files():
    products = {}
    errors = []

    for file_path in sorted(INPUT_DIR.glob("*.xlsx")):
        if file_path.name.startswith("~$"):
            continue

        try:
            wb = load_workbook(
                file_path,
                read_only=True,
                data_only=True
            )
        except Exception as exc:
            errors.append([
                file_path.name,
                0,
                "",
                f"Не удалось открыть файл: {type(exc).__name__}"
            ])
            continue

        try:
            ws = wb.active
            rows = ws.iter_rows(values_only=True)
            headers = next(rows, None)

            if headers is None:
                errors.append([file_path.name, 0, "", "Пустой файл"])
                continue

            columns = normalize_headers(headers)

            if not {"article", "name", "price"} <= columns.keys():
                errors.append([
                    file_path.name, 1, "",
                    "Не найдены обязательные столбцы"
                ])
                continue

            for row_number, row in enumerate(rows, start=2):
                if all(value is None for value in row):
                    continue

                article = str(row[columns["article"]] or "").strip()
                name = str(row[columns["name"]] or "").strip()
                price = parse_price(row[columns["price"]])

                if not article or not name or price is None:
                    errors.append([
                        file_path.name,
                        row_number,
                        article,
                        "Отсутствует или некорректна цена/артикул/название"
                    ])
                    continue

                if article not in products or price < products[article]["price"]:
                    products[article] = {
                        "name": name,
                        "price": price,
                    }

        finally:
            wb.close()

    return products, errors


def save_result(products, errors):
    OUTPUT_DIR.mkdir(exist_ok=True)

    wb = Workbook()
    ws_products = wb.active
    ws_products.title = "Товары"
    ws_errors = wb.create_sheet("Ошибки")

    ws_products.append(["Артикул", "Название", "Цена"])
    ws_errors.append(["Файл", "Строка", "Артикул", "Ошибка"])

    for article, product in sorted(products.items()):
        ws_products.append([
            article,
            product["name"],
            float(product["price"]),
        ])

    for error in errors:
        ws_errors.append(error)

    result_path = OUTPUT_DIR / "result.xlsx"
    wb.save(result_path)

    print(f"Обработано товаров: {len(products)}")
    print(f"Найдено ошибок: {len(errors)}")
    print(f"Результат: {result_path}")


if __name__ == "__main__":
    products, errors = process_files()
    save_result(products, errors)
